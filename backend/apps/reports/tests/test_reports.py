"""Commercial and stock reports (P7-T02, P7-T03, D-113)."""

import datetime as dt
from decimal import Decimal
from io import BytesIO

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.catalog.tests.factories import ProductFactory
from apps.closing import monthly
from apps.core.models import Role
from apps.core.tests.factories import MembershipFactory, UserFactory
from apps.fleet import services as fleet
from apps.inventory import services as stock
from apps.inventory.models import LossReason, MovementKind
from apps.inventory.tests.factories import BatchFactory, StockLocationFactory
from apps.partners.tests.factories import CustomerFactory
from apps.reports import selectors
from apps.sales import services as sales
from apps.sales.models import CreditNoteKind, CreditNoteStock

pytestmark = pytest.mark.django_db
D = Decimal
TODAY = timezone.localdate()


def moment(day: dt.date, hour: int = 10) -> dt.datetime:
    """An aware timestamp on `day`, never later than now: a sale dated today at the fixed hour `10` would
    be rejected as "in the future" when the suite runs before 10am local time."""
    when = timezone.make_aware(dt.datetime.combine(day, dt.time(hour)))
    now = timezone.now()
    return when if when <= now else now


@pytest.fixture
def owner():
    return MembershipFactory.create(role=Role.OWNER)


@pytest.fixture
def organization(owner):
    return owner.organization


@pytest.fixture
def actor(owner):
    return owner.user


def client_for(membership):
    client = APIClient()
    client.force_authenticate(membership.user)
    return client


@pytest.fixture
def cold_room(organization):
    return StockLocationFactory.create(organization=organization, name="Chambre froide 1", counter_sales=True)


def receive(organization, location, product, kg: str, *, cost: str = "20.000", day: dt.date = TODAY):
    batch = BatchFactory.create(organization=organization, product=product, unit_cost=D(cost), received_on=day)
    stock.post_movement(
        organization=organization,
        actor=None,
        kind=MovementKind.RECEIPT,
        batch=batch,
        quantity=kg,
        to_location=location,
        occurred_at=moment(day, 8),
    )
    return batch


def sell(organization, actor, location, product, kg: str, price: str, *, customer=None, day: dt.date = TODAY):
    return sales.create_sale(
        organization=organization,
        actor=actor,
        customer=customer or CustomerFactory.create(organization=organization),
        location=location,
        lines=[{"product": product, "quantity": kg, "unit_price": price}],
        sold_at=moment(day),
        payment={"paid_in_full": True, "payment_method": "cash"},
        confirm=True,
    )


@pytest.fixture
def dorade(organization):
    return ProductFactory.create(organization=organization, name="Dorade")


@pytest.fixture
def crevette(organization):
    return ProductFactory.create(organization=organization, name="Crevette")


def test_sales_by_customer_and_by_product_add_up_to_the_closing(organization, actor, cold_room, dorade, crevette):
    receive(organization, cold_room, dorade, "100.000")
    receive(organization, cold_room, crevette, "50.000", cost="30.000")
    corail = CustomerFactory.create(organization=organization, name="Le Corail")
    sell(organization, actor, cold_room, dorade, "10.000", "32.500", customer=corail)
    sell(organization, actor, cold_room, crevette, "4.000", "45.000", customer=corail)
    sell(organization, actor, cold_room, dorade, "6.000", "30.000")

    by_customer = selectors.sales_report(organization, group_by="customer", date_from=TODAY, date_to=TODAY)
    by_product = selectors.sales_report(organization, group_by="product", date_from=TODAY, date_to=TODAY)

    first = by_customer.rows[0]  # biggest revenue first
    assert (first.name, first.sales, first.quantity, first.revenue, first.cogs, first.margin) == (
        "Le Corail",
        2,
        D("14.000"),
        D("505.000"),
        D("320.000"),
        D("185.000"),
    )
    assert {row.name: row.revenue for row in by_product.rows} == {"Dorade": D("505.000"), "Crevette": D("180.000")}
    closing = monthly.month_report(organization, month=TODAY.replace(day=1))
    for report in (by_customer, by_product):
        assert report.totals.revenue == closing.revenue == D("685.000")
        assert report.totals.cogs == closing.cogs == D("440.000")
        assert report.totals.quantity == D("20.000")
        assert report.totals.sales == closing.sales_count == 3


def test_credit_notes_come_off_the_period_they_are_issued_in(organization, actor, cold_room, dorade):
    receive(organization, cold_room, dorade, "100.000")
    sale = sell(organization, actor, cold_room, dorade, "10.000", "40.000")
    invoice = sales.create_invoice(sale=sale, actor=actor, issue=True)
    sales.create_credit_note(
        invoice=invoice,
        actor=actor,
        kind=CreditNoteKind.PARTIAL,
        reason="Retour",
        lines=[{"invoice_line": invoice.lines.get(), "quantity": "2"}],
        stock=CreditNoteStock.RESTOCK,
    )

    report = selectors.sales_report(organization, group_by="product", date_from=TODAY, date_to=TODAY)

    row = report.rows[0]
    assert (row.sales, row.quantity, row.revenue, row.cogs) == (1, D("8.000"), D("320.000"), D("160.000"))
    assert report.totals.revenue == monthly.month_report(organization, month=TODAY.replace(day=1)).revenue


def test_sales_by_period_split_days_weeks_and_months(organization, actor, cold_room, dorade):
    receive(organization, cold_room, dorade, "100.000", day=TODAY - dt.timedelta(days=40))
    old, recent = TODAY - dt.timedelta(days=35), TODAY
    sell(organization, actor, cold_room, dorade, "5.000", "30.000", day=old)
    sell(organization, actor, cold_room, dorade, "5.000", "30.000", day=recent)
    start = old - dt.timedelta(days=1)

    daily = selectors.sales_report(organization, group_by="day", date_from=start, date_to=TODAY)
    weekly = selectors.sales_report(organization, group_by="week", date_from=start, date_to=TODAY)
    monthly_rows = selectors.sales_report(organization, group_by="month", date_from=start, date_to=TODAY)

    assert [row.key for row in daily.rows] == [old.isoformat(), recent.isoformat()]
    assert [row.key for row in weekly.rows] == [
        (old - dt.timedelta(days=old.weekday())).isoformat(),
        (recent - dt.timedelta(days=recent.weekday())).isoformat(),
    ]
    assert [row.key for row in monthly_rows.rows] == sorted({old.strftime("%Y-%m"), recent.strftime("%Y-%m")})
    assert daily.totals.revenue == weekly.totals.revenue == monthly_rows.totals.revenue == D("300.000")


def test_sales_by_truck_and_driver(organization, actor, dorade):
    driver = UserFactory.create(first_name="Sami", last_name="Ben Ali")
    truck = fleet.create_truck(organization=organization, actor=actor, name="Camion 1", plate="123 TUN 4567")
    fleet.open_driver_day(organization=organization, actor=driver, truck=truck, opening_cash=D("50.000"), date=TODAY)
    receive(organization, truck.location, dorade, "50.000")
    counter = StockLocationFactory.create(organization=organization, name="Comptoir", counter_sales=True)
    receive(organization, counter, dorade, "50.000")
    sell(organization, actor, truck.location, dorade, "10.000", "30.000")
    sell(organization, actor, counter, dorade, "5.000", "30.000")

    by_truck = selectors.sales_report(organization, group_by="truck", date_from=TODAY, date_to=TODAY)
    by_driver = selectors.sales_report(organization, group_by="driver", date_from=TODAY, date_to=TODAY)

    assert {row.name: row.revenue for row in by_truck.rows} == {"Camion 1": D("300.000"), "Comptoir": D("150.000")}
    assert {row.name: row.revenue for row in by_driver.rows} == {"Sami Ben Ali": D("300.000"), "—": D("150.000")}


def test_stock_value_is_by_place_and_product_at_purchase_cost(organization, actor, cold_room, dorade, crevette):
    truck = fleet.create_truck(organization=organization, actor=actor, name="Camion 1", plate="123 TUN 4567")
    receive(organization, cold_room, dorade, "100.000", cost="20.000")
    receive(organization, truck.location, crevette, "10.000", cost="30.000")
    sell(organization, actor, cold_room, dorade, "40.000", "32.500")

    rows = selectors.stock_value(organization)

    assert [(r.location_name, r.product_name, r.quantity, r.value) for r in rows] == [
        ("Camion 1", "Crevette", D("10.000"), D("300.000")),
        ("Chambre froide 1", "Dorade", D("60.000"), D("1200.000")),
    ]


def test_slow_moving_lots_are_those_untouched_for_the_given_days(organization, actor, cold_room, dorade, crevette):
    stale = receive(organization, cold_room, dorade, "20.000", day=TODAY - dt.timedelta(days=45))
    fresh = receive(organization, cold_room, crevette, "10.000", day=TODAY - dt.timedelta(days=3))
    sold_out = receive(organization, cold_room, dorade, "5.000", day=TODAY - dt.timedelta(days=60))
    sell(organization, actor, cold_room, dorade, "5.000", "30.000", day=TODAY - dt.timedelta(days=59))

    rows = selectors.slow_moving(organization, days=30, today=TODAY)

    assert [(r.batch_id, r.days_idle, r.quantity, r.value) for r in rows] == [(stale.pk, 45, D("20.000"), D("400.000"))]
    assert fresh.pk not in {r.batch_id for r in rows}
    assert sold_out.pk not in {r.batch_id for r in rows}  # nothing left of it
    assert {r.batch_id for r in selectors.slow_moving(organization, days=2, today=TODAY)} == {stale.pk, fresh.pk}


def test_losses_and_weight_differences_by_product(organization, actor, cold_room, dorade, crevette):
    receive(organization, cold_room, dorade, "100.000")
    receive(organization, cold_room, crevette, "50.000", cost="30.000")
    stock.record_loss(
        organization=organization,
        actor=actor,
        location=cold_room,
        product=dorade,
        quantity="5.000",
        reason=LossReason.SPOILED,
        lost_on=TODAY,
        may_approve=True,
    )

    rows = selectors.losses(organization, date_from=TODAY, date_to=TODAY)

    assert [(r.product_name, r.lost_quantity, r.lost_value) for r in rows] == [("Dorade", D("5.000"), D("100.000"))]
    assert rows[0].difference_quantity == D("0.000")


def test_reports_are_for_the_office_only(organization, owner):
    urls = [
        reverse(name)
        for name in ("report-sales", "report-stock-value", "report-slow-moving", "report-losses", "report-sales-export")
    ]
    allowed = [owner, MembershipFactory.create(organization=organization, role=Role.ACCOUNTANT)]
    allowed.append(MembershipFactory.create(organization=organization, role=Role.MANAGER))
    driver = MembershipFactory.create(organization=organization, role=Role.DRIVER)

    for url in urls:
        assert all(client_for(member).get(url).status_code == 200 for member in allowed), url
        assert client_for(driver).get(url).status_code == 403, url
        assert APIClient().get(url).status_code in (401, 403), url


def test_a_report_never_shows_another_organizations_sales(organization, actor, cold_room, dorade, owner):
    receive(organization, cold_room, dorade, "50.000")
    sell(organization, actor, cold_room, dorade, "10.000", "30.000")
    other = client_for(MembershipFactory.create(role=Role.OWNER))

    body = other.get(reverse("report-sales"), {"group_by": "customer"}).json()
    mine = client_for(owner).get(reverse("report-sales"), {"group_by": "customer"}).json()

    assert body["rows"] == []
    assert body["totals"]["revenue"] == "0.000"
    assert mine["totals"]["revenue"] == "300.000"


def test_the_sales_api_validates_its_filters(owner):
    api = client_for(owner)
    url = reverse("report-sales")

    assert api.get(url, {"group_by": "weather"}).status_code == 400
    assert api.get(url, {"date_from": "2026-05-10", "date_to": "2026-05-01"}).status_code == 400
    assert api.get(url, {"date_from": "not-a-date"}).status_code == 400


def test_the_sales_export_is_an_excel_file_with_a_totals_row(organization, actor, cold_room, dorade, owner):
    receive(organization, cold_room, dorade, "50.000")
    sell(organization, actor, cold_room, dorade, "10.000", "30.000")

    response = client_for(owner).get(reverse("report-sales-export"), {"group_by": "product"})

    assert response.status_code == 200
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    assert ".xlsx" in response["Content-Disposition"]
    sheet = load_workbook(BytesIO(response.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:3] == ("Libellé", "Ventes", "Kg")
    assert rows[1][0] == "Dorade"
    assert D(str(rows[1][3])) == D("300")
    assert rows[-1][0] == "Total"
    assert D(str(rows[-1][3])) == D("300")


def test_every_other_export_downloads(owner, organization, cold_room, dorade):
    receive(organization, cold_room, dorade, "50.000", day=TODAY - dt.timedelta(days=40))
    api = client_for(owner)

    for name in ("report-stock-value-export", "report-slow-moving-export", "report-losses-export"):
        response = api.get(reverse(name))
        assert response.status_code == 200, name
        assert load_workbook(BytesIO(response.content)).active.max_row >= 2, name
