"""One-sheet Excel files: a header row, the rows and a totals row (D-113). Used by the reports and the lists.

Amounts and kilos stay exact numbers (Decimal) formatted to three decimals, so the accountant can sum them.
"""

from collections.abc import Sequence
from decimal import Decimal
from io import BytesIO
from typing import Any

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MONEY_FORMAT = "#,##0.000"


def workbook_bytes(
    *, title: str, headers: Sequence[str], rows: Sequence[Sequence[Any]], totals: Sequence[Any] | None = None
) -> bytes:
    sheet = Workbook().active
    assert sheet is not None
    sheet.title = title[:31]
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row in rows:
        sheet.append(list(row))
    if totals is not None:
        sheet.append(list(totals))
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
    # Names come from people (a driver's own display name, customers, suppliers, products): openpyxl would store a
    # text starting with "=" as a formula, which Excel runs when the owner opens the file (D-141). Text stays text.
    for line in sheet.iter_rows():
        for cell in line:
            if isinstance(cell.value, str) and cell.data_type == "f":
                cell.data_type = "s"
    for column in sheet.iter_cols(min_row=1):
        width = max((len(str(cell.value)) for cell in column if cell.value is not None), default=8)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(width + 2, 10), 40)
        for cell in column[1:]:
            if isinstance(cell.value, Decimal):
                cell.number_format = MONEY_FORMAT
    sheet.freeze_panes = "A2"
    out = BytesIO()
    sheet.parent.save(out)
    return out.getvalue()


def xlsx_response(name: str, content: bytes) -> HttpResponse:
    """The file as a download named `<name>.xlsx`."""
    response = HttpResponse(content, content_type=XLSX_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{name}.xlsx"'
    return response
