import datetime as dt
import io
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.catalog import selectors as catalog
from apps.catalog.tests.factories import CategoryFactory
from apps.core.models import AuditLog, Role
from apps.core.tests.factories import MembershipFactory
from apps.dataimport import services
from apps.dataimport.columns import KINDS
from apps.hr import selectors as hr
from apps.partners import selectors as partners

pytestmark = pytest.mark.django_db


@pytest.fixture
def owner():
    return MembershipFactory.create(role=Role.OWNER)


@pytest.fixture
def organization(owner):
    return owner.organization


def client_for(membership):
    client = APIClient()
    client.force_authenticate(membership.user)
    return client


def sheet(kind, rows, *, titles=None):
    """A filled template: the kind's titles (or `titles`), then one row per dict of column key → value."""
    book = Workbook()
    active = book.active
    columns = KINDS[kind].columns
    active.append(titles or [column.title for column in columns])
    for row in rows:
        active.append([row.get(column.key) for column in columns])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def run(organization, owner, kind, rows, *, dry_run=False, **options):
    return services.run_import(
        organization=organization,
        actor=owner.user,
        kind_key=kind,
        content=sheet(kind, rows, **options),
        dry_run=dry_run,
    )


# ---- templates ----


@pytest.mark.parametrize("kind", list(KINDS))
def test_every_template_opens_in_excel_with_its_titles_and_a_help_sheet(owner, kind):
    response = client_for(owner).get(reverse("import-template", args=[kind]))

    assert response.status_code == 200
    assert response["Content-Disposition"] == f'attachment; filename="modele-{kind}.xlsx"'
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == [KINDS[kind].title, "Aide"]
    titles = [cell.value for cell in book.worksheets[0][1]]
    assert titles[0].startswith(KINDS[kind].columns[0].title)
    assert len(titles) == len(KINDS[kind].columns)


def test_the_filled_template_reads_back_through_the_titles(organization, owner):
    # The downloaded template, typed in, is what the import expects.
    _name, content = services.template_for("suppliers")
    book = load_workbook(io.BytesIO(content))
    book.worksheets[0].append(["F-100", "Pêcherie du Golfe"])
    output = io.BytesIO()
    book.save(output)

    report = services.run_import(
        organization=organization, actor=owner.user, kind_key="suppliers", content=output.getvalue(), dry_run=False
    )

    assert (report.created, report.error_count, report.committed) == (1, 0, True)


# ---- suppliers, customers, employees ----


def test_suppliers_are_created_through_the_service_and_audited(organization, owner):
    report = run(
        organization,
        owner,
        "suppliers",
        [
            {"reference": "F-001", "name": "Pêcherie du Golfe", "tax_id": "1234567A/A/M/000", "city": "Sfax"},
            {"name": "Sans référence"},
        ],
    )

    assert (report.created, report.skipped, report.error_count, report.committed) == (2, 0, 0, True)
    names = {row.reference: row.name for row in partners.suppliers_of(organization)}
    assert names["F-001"] == "Pêcherie du Golfe"
    # A row without a reference gets the next free one.
    assert "Sans référence" in names.values()
    assert AuditLog.objects.filter(organization=organization, action="supplier.created").count() == 2


def test_a_dry_run_reports_everything_and_writes_nothing(organization, owner):
    rows = [{"reference": "F-001", "name": "Pêcherie du Golfe"}, {"reference": "F-002", "name": ""}]

    report = run(organization, owner, "suppliers", rows, dry_run=True)

    assert (report.created, report.error_count, report.committed) == (1, 1, False)
    assert partners.suppliers_of(organization, status="all").count() == 0
    assert not AuditLog.objects.filter(organization=organization, action="supplier.created").exists()


def test_one_error_anywhere_writes_nothing_even_for_a_real_import(organization, owner):
    rows = [{"reference": "F-001", "name": "Bon"}, {"reference": "F-002", "email": "pas-un-mail", "name": "Mauvais"}]

    report = run(organization, owner, "suppliers", rows)

    assert (report.created, report.error_count, report.committed) == (1, 1, False)
    bad = report.rows[1]
    assert (bad.row, bad.status) == (3, "error")
    assert [error.column for error in bad.errors] == ["E-mail"]
    assert partners.suppliers_of(organization, status="all").count() == 0


def test_a_known_reference_is_skipped_so_a_file_can_be_sent_again(organization, owner):
    rows = [{"reference": "F-001", "name": "Pêcherie du Golfe"}]
    run(organization, owner, "suppliers", rows)

    again = run(organization, owner, "suppliers", [*rows, {"reference": "F-002", "name": "Nouveau"}])

    assert (again.created, again.skipped, again.error_count) == (1, 1, 0)
    assert partners.suppliers_of(organization).count() == 2


def test_a_reference_twice_in_the_file_is_an_error_on_the_second(organization, owner):
    report = run(
        organization, owner, "suppliers", [{"reference": "F-1", "name": "A"}, {"reference": "f-1", "name": "B"}]
    )

    assert [row.status for row in report.rows] == ["created", "error"]
    assert "plus haut" in report.rows[1].errors[0].message


def test_customers_take_a_kind_by_label_or_code_and_a_credit_limit_with_a_comma(organization, owner):
    report = run(
        organization,
        owner,
        "customers",
        [
            {"name": "Restaurant Le Corail", "kind": "Restaurant", "credit_limit": "500,500", "payment_terms_days": 30},
            {"name": "Hôtel Mouradi", "kind": "hotel", "credit_limit": 2000},
        ],
    )

    assert (report.created, report.error_count) == (2, 0)
    by_name = {row.name: row for row in partners.customers_of(organization)}
    assert by_name["Restaurant Le Corail"].credit_limit == Decimal("500.500")
    assert by_name["Restaurant Le Corail"].payment_terms_days == 30
    assert by_name["Hôtel Mouradi"].kind == "hotel"


def test_a_customer_with_an_unknown_kind_or_a_bad_number_is_reported_per_column(organization, owner):
    report = run(
        organization,
        owner,
        "customers",
        [{"name": "A", "kind": "Cantine"}, {"name": "B", "kind": "hotel", "credit_limit": "beaucoup"}],
        dry_run=True,
    )

    first, second = report.rows
    assert [(error.column, "Type inconnu" in error.message) for error in first.errors] == [("Type", True)]
    assert [error.column for error in second.errors] == ["Limite de crédit (DT)"]


def test_employees_take_a_hire_date_typed_in_french_and_a_salary(organization, owner):
    report = run(
        organization,
        owner,
        "employees",
        [
            {
                "name": "Ali Ben Salah",
                "position": "Manutentionnaire",
                "hire_date": "01/03/2024",
                "monthly_salary": "1200,000",
            },
            {"name": "Sami", "hire_date": dt.date(2025, 1, 15), "monthly_salary": 900},
        ],
    )

    assert (report.created, report.error_count) == (2, 0)
    by_name = {row.name: row for row in hr.employees_of(organization)}
    assert by_name["Ali Ben Salah"].hire_date == dt.date(2024, 3, 1)
    assert by_name["Ali Ben Salah"].monthly_salary == Decimal("1200.000")
    assert by_name["Sami"].hire_date == dt.date(2025, 1, 15)


# ---- products ----


def test_products_need_an_existing_category_and_a_price(organization, owner):
    CategoryFactory.create(organization=organization, name="Crustacés")
    report = run(
        organization,
        owner,
        "products",
        [
            {"reference": "CR-001", "name": "Crevette Royale", "category": "crustacés", "sale_price": "42,000"},
            {"reference": "XX-1", "name": "Mystère", "category": "Inconnue", "sale_price": 10},
            {"reference": "XX-2", "name": "Gratuit", "category": "Crustacés", "sale_price": 0},
        ],
        dry_run=True,
    )

    assert [row.status for row in report.rows] == ["created", "error", "error"]
    assert "Catégories existantes : Crustacés" in report.rows[1].errors[0].message
    assert report.rows[1].errors[0].column == "Catégorie"
    assert report.rows[2].errors[0].column == "Prix de vente (DT/kg)"


def test_a_real_product_import_creates_the_product_with_its_first_price(organization, owner):
    CategoryFactory.create(organization=organization, name="Crustacés")

    report = run(
        organization,
        owner,
        "products",
        [
            {
                "reference": "CR-001",
                "name": "Crevette Royale",
                "category": "Crustacés",
                "sale_price": "42,000",
                "vat_rate": 0,
            }
        ],
    )

    assert (report.created, report.committed) == (1, True)
    (product,) = catalog.products_of(organization)
    price = product.current_price  # type: ignore[attr-defined]
    assert (product.reference, price, product.vat_rate) == ("CR-001", Decimal("42.000"), Decimal("0"))


# ---- the file itself ----


def test_a_missing_required_column_refuses_the_whole_file(organization, owner):
    with pytest.raises(services.WorkbookError, match="Nom"):
        run(organization, owner, "suppliers", [{"reference": "F-1"}], titles=["Référence", "Ville"])


def test_a_file_that_is_not_excel_is_refused(organization, owner):
    with pytest.raises(services.WorkbookError, match="Excel"):
        services.run_import(
            organization=organization, actor=owner.user, kind_key="suppliers", content=b"nom;ville", dry_run=True
        )


def test_titles_are_matched_without_case_accents_or_the_star(organization, owner):
    report = run(
        organization,
        owner,
        "suppliers",
        [{"reference": "F-1", "name": "Pêcherie"}],
        titles=["REFERENCE", "nom *", "Ville"],
    )

    assert (report.created, report.error_count) == (1, 0)


# ---- the API ----


def upload(api, kind, content, **data):
    file = SimpleUploadedFile("modele.xlsx", content, content_type="application/octet-stream")
    return api.post(reverse("import-run", args=[kind]), {"file": file, **data}, format="multipart")


def test_the_owner_uploads_a_dry_run_then_confirms(organization, owner):
    api = client_for(owner)
    content = sheet("suppliers", [{"reference": "F-001", "name": "Pêcherie du Golfe"}])

    checked = upload(api, "suppliers", content)

    assert checked.status_code == 200
    body = checked.json()
    assert (body["dry_run"], body["committed"], body["created"], body["error_count"]) == (True, False, 1, 0)
    assert body["rows"][0] == {
        "row": 2,
        "status": "created",
        "reference": "F-001",
        "name": "Pêcherie du Golfe",
        "errors": [],
    }
    assert partners.suppliers_of(organization, status="all").count() == 0

    done = upload(api, "suppliers", content, dry_run="false")

    assert done.json()["committed"] is True
    assert partners.suppliers_of(organization).count() == 1


def test_the_report_names_the_column_and_the_reason_of_each_error(organization, owner):
    content = sheet("suppliers", [{"reference": "F-001", "name": "A", "email": "oups"}])

    body = upload(client_for(owner), "suppliers", content).json()

    assert body["error_count"] == 1
    assert body["rows"][0]["errors"] == [{"column": "E-mail", "message": "Adresse e-mail invalide."}]


@pytest.mark.parametrize("role", [Role.MANAGER, Role.ACCOUNTANT, Role.DRIVER])
def test_only_the_owner_may_download_a_template_or_import(role):
    api = client_for(MembershipFactory.create(role=role))
    content = sheet("suppliers", [{"name": "A"}])

    assert api.get(reverse("import-template", args=["suppliers"])).status_code == 403
    assert upload(api, "suppliers", content, dry_run="false").status_code == 403


def test_a_stranger_gets_401_and_an_unknown_kind_404(owner):
    assert APIClient().get(reverse("import-template", args=["suppliers"])).status_code in (401, 403)
    assert client_for(owner).get(reverse("import-template", args=["lutins"])).status_code == 404


def test_an_unreadable_file_answers_400_on_the_file_field(owner):
    response = upload(client_for(owner), "suppliers", b"pas un classeur")

    assert response.status_code == 400
    assert "Excel" in response.json()["file"][0]


# ---- opening stock ----


def opening_setup(organization):
    from apps.catalog.tests.factories import ProductFactory
    from apps.inventory.tests.factories import StockLocationFactory
    from apps.partners.tests.factories import SupplierFactory

    ProductFactory.create(organization=organization, reference="CR-001", name="Crevette")
    SupplierFactory.create(organization=organization, reference="F-001", name="Golfe")
    return StockLocationFactory.create(organization=organization, name="Chambre froide 1")


OPENING = {
    "product": "CR-001",
    "location": "Chambre froide 1",
    "supplier": "F-001",
    "quantity": "120,500",
    "unit_cost": "21,000",
    "received_on": "05/10/2026",
}


def test_opening_stock_posts_a_receipt_into_the_location(organization, owner):
    from apps.inventory import selectors as inventory

    location = opening_setup(organization)

    report = run(organization, owner, "opening_stock", [OPENING])

    assert (report.created, report.error_count, report.committed) == (1, 0, True)
    (batch,) = inventory.batches_of(organization)
    assert (batch.unit_cost, batch.received_on, batch.receipt_reference) == (
        Decimal("21.000"),
        dt.date(2026, 10, 5),
        "STOCK-OUVERTURE",
    )
    assert inventory.quantity_on_hand(organization, location_id=location.pk, product_id=batch.product_id) == Decimal(
        "120.500"
    )


def test_opening_stock_sent_twice_does_not_double_the_stock(organization, owner):
    from apps.inventory import selectors as inventory

    opening_setup(organization)
    run(organization, owner, "opening_stock", [OPENING])

    again = run(organization, owner, "opening_stock", [OPENING])

    assert (again.created, again.skipped) == (0, 1)
    assert inventory.batches_of(organization).count() == 1


def test_opening_stock_names_the_unknown_product_location_and_supplier(organization, owner):
    opening_setup(organization)
    bad = {**OPENING, "product": "ZZ", "location": "Cave", "supplier": "F-999"}

    report = run(organization, owner, "opening_stock", [bad, {**OPENING, "quantity": 0}], dry_run=True)

    first, second = report.rows
    assert {error.column for error in first.errors} == {"Produit (référence)", "Emplacement", "Fournisseur (référence)"}
    assert [error.column for error in second.errors] == ["Quantité (kg)"]


def test_the_same_lot_twice_in_the_file_is_an_error(organization, owner):
    opening_setup(organization)

    report = run(organization, owner, "opening_stock", [OPENING, OPENING], dry_run=True)

    assert [row.status for row in report.rows] == ["created", "error"]


# ---- opening debts ----


def test_customer_debts_become_receivables_that_age_from_their_date(organization, owner):
    from apps.partners.tests.factories import CustomerFactory
    from apps.sales import receivables

    customer = CustomerFactory.create(organization=organization, reference="C-001", name="Corail")
    old = (timezone.localdate() - dt.timedelta(days=45)).strftime("%d/%m/%Y")

    report = run(
        organization,
        owner,
        "customer_debts",
        [{"customer": "C-001", "amount": "1250,000", "owed_on": old, "notes": "Facture papier 118"}],
    )

    assert (report.created, report.error_count, report.committed) == (1, 0, True)
    balance = receivables.customer_balance(organization, customer.pk)
    assert (balance.balance, balance.buckets["d31_60"]) == (Decimal("1250.000"), Decimal("1250.000"))
    (open_sale,) = balance.open_sales
    assert open_sale.number.startswith("OUV-")


def test_an_opening_receivable_is_never_invoiced_nor_counted_as_a_sale(organization, owner):
    from django.core.exceptions import ValidationError

    from apps.partners.tests.factories import CustomerFactory
    from apps.sales import selectors
    from apps.sales import services as sales

    CustomerFactory.create(organization=organization, reference="C-001", name="Corail")
    run(organization, owner, "customer_debts", [{"customer": "C-001", "amount": 100, "owed_on": "01/09/2026"}])
    (sale,) = selectors.sales_of(organization)

    assert sale.is_opening
    assert not list(selectors.sales_of(organization, uninvoiced=True))
    with pytest.raises(ValidationError, match="ouverture"):
        sales.create_invoice(apply_vat=True, sale=sale, actor=owner.user)


def test_customer_debts_sent_twice_are_skipped_and_unknown_customers_reported(organization, owner):
    from apps.partners.tests.factories import CustomerFactory

    CustomerFactory.create(organization=organization, reference="C-001", name="Corail")
    row = {"customer": "C-001", "amount": 100, "owed_on": "01/09/2026"}
    run(organization, owner, "customer_debts", [row])

    again = run(organization, owner, "customer_debts", [row], dry_run=True)
    unknown = run(organization, owner, "customer_debts", [{**row, "customer": "C-999"}], dry_run=True)

    assert (again.created, again.skipped) == (0, 1)
    assert unknown.rows[0].errors[0].column == "Client (référence)"


def test_supplier_debts_become_posted_payables(organization, owner):
    from apps.partners.tests.factories import SupplierFactory
    from apps.purchasing import selectors as purchasing

    supplier = SupplierFactory.create(organization=organization, reference="F-001", name="Golfe")

    report = run(
        organization,
        owner,
        "supplier_debts",
        [{"supplier": "F-001", "amount": "3400,000", "invoice_date": "20/09/2026", "due_date": "20/10/2026"}],
    )

    assert (report.created, report.error_count) == (1, 0)
    (invoice,) = purchasing.invoices_of(organization, supplier_id=supplier.pk)
    assert (invoice.status, invoice.total, invoice.is_opening) == ("posted", Decimal("3400.000"), True)
    assert purchasing.supplier_balance(organization, supplier.pk).outstanding == Decimal("3400.000")

    again = run(
        organization,
        owner,
        "supplier_debts",
        [{"supplier": "F-001", "amount": "3400,000", "invoice_date": "20/09/2026"}],
    )
    assert again.skipped == 1


def test_a_debt_row_with_a_bad_amount_or_a_future_date_is_an_error(organization, owner):
    from apps.partners.tests.factories import CustomerFactory

    CustomerFactory.create(organization=organization, reference="C-001", name="Corail")
    future = (timezone.localdate() + dt.timedelta(days=5)).strftime("%d/%m/%Y")

    report = run(
        organization,
        owner,
        "customer_debts",
        [
            {"customer": "C-001", "amount": 0, "owed_on": "01/09/2026"},
            {"customer": "C-001", "amount": 10, "owed_on": future},
        ],
        dry_run=True,
    )

    assert [row.status for row in report.rows] == ["error", "error"]
